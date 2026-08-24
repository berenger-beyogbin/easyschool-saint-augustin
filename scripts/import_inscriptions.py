"""Import sécurisé de l'ancien export Excel INSCRIPTION dans TInscription.

L'ancien export référence les élèves/familles par leur ancien ID (préservé
lors de l'import de TFamille/Eleve) et la classe par un ancien IdClasse, qui
est traduit vers la nouvelle TClasse via l'export NIVEAU CLASSE (table de
correspondance ancien IdClasse -> libellé -> TClasse de la nouvelle base).

Sans --commit, la transaction est toujours annulée après simulation.

Exemple :
  python scripts/import_inscriptions.py "E:\\EXPORT EPC\\Export INSCRIPTION.xlsx" "E:\\EXPORT EPC\\Export NIVEAU CLASSE.xlsx" --annee "2026-2027" --analyze-only
  python scripts/import_inscriptions.py "E:\\EXPORT EPC\\Export INSCRIPTION.xlsx" "E:\\EXPORT EPC\\Export NIVEAU CLASSE.xlsx" --annee "2026-2027" --commit
"""

from __future__ import annotations

import argparse
import csv
import sys
import tempfile
import zipfile
from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

PROJECT_ROOT = Path(__file__).resolve().parents[1]
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))

INSCRIPTION_HEADERS = [
    "IDTInscription", "IDTAnneeScolaire", "IDFamille", "IDEleve", "IDNiveau",
    "IdClasse", "Nouveau", "Scolarite", "Transport", "Cantine", "AutresFrais",
    "Login", "DateInscrip",
]
NIVEAU_CLASSE_HEADERS = [
    "IDNiveau", "Niv_Libelle", "IDPNiveau", "SINiveau",
]


@dataclass
class Rejection:
    excel_row: int
    legacy_id: str
    reason: str


def clean(value: Any) -> str:
    return "" if value is None else str(value).strip()


def as_int(value: Any) -> int | None:
    raw = clean(value)
    if not raw:
        return None
    try:
        return int(float(raw.replace(",", ".")))
    except ValueError:
        return None


def as_bool(value: Any) -> bool:
    return as_int(value) == 1


def as_date(value: Any) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    raw = clean(value)
    for fmt in ("%d/%m/%Y", "%Y-%m-%d", "%d-%m-%Y"):
        try:
            return datetime.strptime(raw, fmt).date()
        except ValueError:
            pass
    return None


def normalize_libelle(value: str) -> str:
    return " ".join(clean(value).upper().split())


def repaired_workbook(source: Path, temp_dir: Path) -> Path:
    """Corrige dans une copie temporaire la faute XML `biltinId` de l'export."""
    repaired = temp_dir / (source.stem + "_repare.xlsx")
    with zipfile.ZipFile(source, "r") as incoming, zipfile.ZipFile(repaired, "w") as outgoing:
        for item in incoming.infolist():
            data = incoming.read(item.filename)
            if item.filename == "xl/styles.xml":
                data = data.replace(b"biltinId", b"builtinId")
            outgoing.writestr(item, data)
    return repaired


def read_rows(path: Path, expected_headers: list[str]) -> list[tuple[int, dict[str, Any]]]:
    with tempfile.TemporaryDirectory(prefix="easy_school_inscriptions_") as folder:
        repaired = repaired_workbook(path, Path(folder))
        workbook = load_workbook(repaired, read_only=True, data_only=True)
        sheet = workbook.active
        iterator = sheet.iter_rows(values_only=True)
        headers = [clean(value) for value in next(iterator)]
        missing = [header for header in expected_headers if header not in headers]
        if missing:
            raise ValueError(f"Colonnes absentes de '{path.name}' : {', '.join(missing)}")
        rows = [(number, dict(zip(headers, values))) for number, values in enumerate(iterator, start=2)]
        workbook.close()
        return rows


def build_classe_lookup(niveau_classe_rows: list[tuple[int, dict[str, Any]]]) -> dict[int, str]:
    """Ancien IdClasse (SINiveau=0, une vraie classe, pas un niveau) -> libellé normalisé."""
    lookup: dict[int, str] = {}
    for _, row in niveau_classe_rows:
        if as_int(row.get("SINiveau")) == 1:
            continue  # C'est un niveau parent, pas une classe
        old_id = as_int(row.get("IDNiveau"))
        libelle = clean(row.get("Niv_Libelle"))
        if old_id is not None and libelle:
            lookup[old_id] = normalize_libelle(libelle)
    return lookup


def write_rejections(path: Path, rows: list[Rejection]) -> None:
    with path.open("w", newline="", encoding="utf-8-sig") as handle:
        writer = csv.writer(handle, delimiter=";")
        writer.writerow(["Ligne Excel", "Ancien IDTInscription", "Motif"])
        writer.writerows((r.excel_row, r.legacy_id, r.reason) for r in rows)


def main() -> int:
    parser = argparse.ArgumentParser(description="Importer l'ancien export INSCRIPTION vers TInscription.")
    parser.add_argument("inscriptions_xlsx", type=Path, help="Chemin de Export INSCRIPTION.xlsx")
    parser.add_argument("niveau_classe_xlsx", type=Path, help="Chemin de Export NIVEAU CLASSE.xlsx")
    parser.add_argument("--annee", required=True, help="Libellé de l'année scolaire cible, ex: 2026-2027")
    parser.add_argument("--commit", action="store_true", help="Valider réellement l'import en base.")
    parser.add_argument("--analyze-only", action="store_true", help="Analyser les exports sans connexion à la base.")
    parser.add_argument("--rejects", type=Path, default=Path("inscriptions_rejetees.csv"))
    args = parser.parse_args()

    if not args.inscriptions_xlsx.is_file():
        parser.error(f"Fichier introuvable : {args.inscriptions_xlsx}")
    if not args.niveau_classe_xlsx.is_file():
        parser.error(f"Fichier introuvable : {args.niveau_classe_xlsx}")

    from sqlalchemy import select, text
    from sqlalchemy.exc import IntegrityError

    from app.database import get_session
    import models  # noqa: F401 - enregistre tous les mappers SQLAlchemy avant usage
    from models.inscription import TInscription
    from models.annee_scolaire import TAnneeScolaire
    from models.classe import TClasse
    from models.eleve import Eleve
    from models.famille import TFamille

    inscription_rows = read_rows(args.inscriptions_xlsx, INSCRIPTION_HEADERS)
    niveau_classe_rows = read_rows(args.niveau_classe_xlsx, NIVEAU_CLASSE_HEADERS)
    old_classe_to_libelle = build_classe_lookup(niveau_classe_rows)

    if args.analyze_only:
        classes_referencees = {
            as_int(row.get("IdClasse")) for _, row in inscription_rows
            if as_int(row.get("IdClasse")) is not None
        }
        classes_absentes = sorted(classes_referencees - old_classe_to_libelle.keys())
        print(f"Inscriptions lues : {len(inscription_rows)}")
        print(f"Classes de correspondance : {len(old_classe_to_libelle)}")
        print(f"Classes référencées sans correspondance : {len(classes_absentes)}")
        if classes_absentes:
            print("Anciens IdClasse absents : " + ", ".join(map(str, classes_absentes)))
        return 0

    session = get_session()
    inserted = 0
    rejected: list[Rejection] = []
    try:
        annee = session.scalar(select(TAnneeScolaire).where(TAnneeScolaire.Libelle == args.annee))
        if annee is None:
            print(f"Année scolaire '{args.annee}' introuvable dans la base.", file=sys.stderr)
            return 2

        classe_by_libelle = {
            normalize_libelle(libelle): (id_classe, id_niveau)
            for id_classe, libelle, id_niveau in session.execute(
                select(TClasse.IDTClasse, TClasse.LibClasse, TClasse.IDT_Niveau)
            )
        }
        eleve_ids = set(session.scalars(select(Eleve.IDEleve)))
        famille_ids = set(session.scalars(select(TFamille.IdTFamille)))
        existing_pairs = {
            (id_eleve, id_annee)
            for id_eleve, id_annee in session.execute(
                select(TInscription.IDEleve, TInscription.IDTAnneeScolaire)
            )
        }

        seen_pairs: set[tuple[int, int]] = set()
        for excel_row, source in inscription_rows:
            legacy_id = clean(source.get("IDTInscription"))
            id_eleve = as_int(source.get("IDEleve"))
            id_famille = as_int(source.get("IDFamille"))
            old_id_classe = as_int(source.get("IdClasse"))
            date_inscription = as_date(source.get("DateInscrip"))

            reason = ""
            if id_eleve is None or id_eleve not in eleve_ids:
                reason = f"Élève {id_eleve} introuvable dans la nouvelle base"
            elif id_famille is None or id_famille not in famille_ids:
                reason = f"Famille {id_famille} introuvable dans la nouvelle base"
            elif old_id_classe is None or old_id_classe not in old_classe_to_libelle:
                reason = f"Ancienne classe {old_id_classe} introuvable dans l'export NIVEAU CLASSE"
            elif not date_inscription:
                reason = "Date d'inscription absente ou invalide"
            else:
                libelle = old_classe_to_libelle[old_id_classe]
                match = classe_by_libelle.get(libelle)
                if match is None:
                    reason = f"Classe '{libelle}' introuvable dans la nouvelle base (TClasse)"

            if reason:
                rejected.append(Rejection(excel_row, legacy_id, reason))
                continue

            id_classe, id_niveau = classe_by_libelle[old_classe_to_libelle[old_id_classe]]
            pair = (id_eleve, annee.IDTAnneeScolaire)
            if pair in existing_pairs or pair in seen_pairs:
                rejected.append(Rejection(excel_row, legacy_id, "Élève déjà inscrit pour cette année"))
                continue
            seen_pairs.add(pair)

            session.add(TInscription(
                IDTAnneeScolaire=annee.IDTAnneeScolaire,
                IDFamille=id_famille,
                IDEleve=id_eleve,
                IDNiveau=id_niveau,
                IDClasse=id_classe,
                Nouveau=as_bool(source.get("Nouveau")),
                Scolarite=as_bool(source.get("Scolarite")),
                Transport=as_bool(source.get("Transport")),
                Cantine=as_bool(source.get("Cantine")),
                AutresFrais=as_bool(source.get("AutresFrais")),
                Login=clean(source.get("Login"))[:50] or None,
                DateInscription=date_inscription,
            ))
            inserted += 1

        session.flush()
        if inserted:
            session.execute(text(
                "SELECT setval(pg_get_serial_sequence('\"TInscription\"', 'IDTInscription'), "
                "COALESCE((SELECT MAX(\"IDTInscription\") FROM \"TInscription\"), 1), true)"
            ))
        write_rejections(args.rejects, rejected)
        if args.commit:
            session.commit()
            action = "IMPORT VALIDÉ"
        else:
            session.rollback()
            action = "SIMULATION : aucune modification enregistrée"
        print(action)
        print(f"Lignes lues : {len(inscription_rows)}")
        print(f"Lignes prêtes à être insérées : {inserted}")
        print(f"Lignes rejetées/ignorées : {len(rejected)}")
        print(f"Rapport : {args.rejects.resolve()}")
        return 0
    except IntegrityError as exc:
        session.rollback()
        print(f"Import annulé à cause d'une contrainte de base : {exc.orig}", file=sys.stderr)
        return 2
    except Exception as exc:
        session.rollback()
        print(f"Import annulé : {exc}", file=sys.stderr)
        return 2
    finally:
        session.close()


if __name__ == "__main__":
    raise SystemExit(main())
