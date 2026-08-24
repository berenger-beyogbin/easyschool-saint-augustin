"""Import sécurisé des anciens exports Excel VERSEMENTSCOL et VERSEMENTCANT
dans la table unifiée VersementScol.

L'ancien système avait deux tables séparées (scolarité+transport d'un côté,
cantine de l'autre). La nouvelle base regroupe tout dans VersementScol avec
4 colonnes de montant. Pour rester fidèle à l'historique, chaque ancienne
ligne devient une nouvelle ligne distincte, avec seulement la colonne
concernée renseignée (les autres à 0) — aucune fusion de paiements distincts.

Les anciens ID de versement ne sont PAS conservés (aucune autre table n'y
fait référence) : de nouveaux ID sont générés à l'insertion.

Sans --commit, la transaction est toujours annulée après simulation.

Exemple :
  python scripts/import_versements.py "E:\\EXPORT EPC\\Export versementscol.xlsx" "E:\\EXPORT EPC\\Export versementcant.xlsx" --annee "2026-2027" --analyze-only
  python scripts/import_versements.py "E:\\EXPORT EPC\\Export versementscol.xlsx" "E:\\EXPORT EPC\\Export versementcant.xlsx" --annee "2026-2027" --commit
"""

from __future__ import annotations

import argparse
import csv
import sys
import tempfile
import zipfile
from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

PROJECT_ROOT = Path(__file__).resolve().parents[1]
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))

VERSEMENTSCOL_HEADERS = [
    "IDVersementScol", "IDFamille", "DateVers", "MontantVersTrans",
    "MontantVersSco", "IDTAnneeScolaire", "IDEleve", "Restitution",
    "Login", "Reduction", "Impaye",
]
VERSEMENTCANT_HEADERS = [
    "IDVersementCant", "IDFamille", "DateVers", "MontantVers",
    "IDTAnneeScolaire", "IDEleve", "Reduction", "Restitution", "Login", "Impaye",
]


@dataclass
class Rejection:
    source: str
    excel_row: int
    legacy_id: str
    reason: str


def clean(value: Any) -> str:
    return "" if value is None else str(value).strip()


def as_int(value: Any) -> int | None:
    raw = clean(value)
    if not raw:
        return None
    try:
        return int(float(raw.replace(",", ".")))
    except ValueError:
        return None


def as_decimal(value: Any) -> float:
    raw = clean(value)
    if not raw:
        return 0.0
    try:
        return float(raw.replace(",", "."))
    except ValueError:
        return 0.0


def as_bool(value: Any) -> bool:
    return as_int(value) == 1


def as_date(value: Any) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    raw = clean(value)
    for fmt in ("%d/%m/%Y", "%Y-%m-%d", "%d-%m-%Y"):
        try:
            return datetime.strptime(raw, fmt).date()
        except ValueError:
            pass
    return None


def repaired_workbook(source: Path, temp_dir: Path) -> Path:
    """Corrige dans une copie temporaire la faute XML `biltinId` de l'export."""
    repaired = temp_dir / (source.stem + "_repare.xlsx")
    with zipfile.ZipFile(source, "r") as incoming, zipfile.ZipFile(repaired, "w") as outgoing:
        for item in incoming.infolist():
            data = incoming.read(item.filename)
            if item.filename == "xl/styles.xml":
                data = data.replace(b"biltinId", b"builtinId")
            outgoing.writestr(item, data)
    return repaired


def read_rows(path: Path, expected_headers: list[str]) -> list[tuple[int, dict[str, Any]]]:
    with tempfile.TemporaryDirectory(prefix="easy_school_versements_") as folder:
        repaired = repaired_workbook(path, Path(folder))
        workbook = load_workbook(repaired, read_only=True, data_only=True)
        sheet = workbook.active
        iterator = sheet.iter_rows(values_only=True)
        headers = [clean(value) for value in next(iterator)]
        missing = [header for header in expected_headers if header not in headers]
        if missing:
            raise ValueError(f"Colonnes absentes de '{path.name}' : {', '.join(missing)}")
        rows = [(number, dict(zip(headers, values))) for number, values in enumerate(iterator, start=2)]
        workbook.close()
        return rows


def write_rejections(path: Path, rows: list[Rejection]) -> None:
    with path.open("w", newline="", encoding="utf-8-sig") as handle:
        writer = csv.writer(handle, delimiter=";")
        writer.writerow(["Source", "Ligne Excel", "Ancien ID", "Motif"])
        writer.writerows((r.source, r.excel_row, r.legacy_id, r.reason) for r in rows)


def main() -> int:
    parser = argparse.ArgumentParser(description="Importer versementscol + versementcant vers VersementScol.")
    parser.add_argument("versementscol_xlsx", type=Path)
    parser.add_argument("versementcant_xlsx", type=Path)
    parser.add_argument("--annee", required=True, help="Libellé de l'année scolaire cible, ex: 2026-2027")
    parser.add_argument("--commit", action="store_true")
    parser.add_argument("--analyze-only", action="store_true", help="Analyser les exports sans connexion à la base.")
    parser.add_argument("--rejects", type=Path, default=Path("versements_rejetes.csv"))
    args = parser.parse_args()

    for path in (args.versementscol_xlsx, args.versementcant_xlsx):
        if not path.is_file():
            parser.error(f"Fichier introuvable : {path}")

    from sqlalchemy import select, text
    from sqlalchemy.exc import IntegrityError

    from app.database import get_session
    import models  # noqa: F401 - enregistre tous les mappers SQLAlchemy avant usage
    from models.versement_scol import VersementScol
    from models.annee_scolaire import TAnneeScolaire
    from models.eleve import Eleve
    from models.famille import TFamille

    scol_rows = read_rows(args.versementscol_xlsx, VERSEMENTSCOL_HEADERS)
    cant_rows = read_rows(args.versementcant_xlsx, VERSEMENTCANT_HEADERS)

    if args.analyze_only:
        invalid_dates = sum(
            as_date(row.get("DateVers")) is None
            for _, row in scol_rows + cant_rows
        )
        print(f"Lignes lues : {len(scol_rows) + len(cant_rows)} (scol={len(scol_rows)}, cant={len(cant_rows)})")
        print(f"Dates absentes ou invalides : {invalid_dates}")
        return 0

    session = get_session()
    inserted = 0
    rejected: list[Rejection] = []
    try:
        annee = session.scalar(select(TAnneeScolaire).where(TAnneeScolaire.Libelle == args.annee))
        if annee is None:
            print(f"Année scolaire '{args.annee}' introuvable dans la base.", file=sys.stderr)
            return 2

        eleve_ids = set(session.scalars(select(Eleve.IDEleve)))
        famille_ids = set(session.scalars(select(TFamille.IdTFamille)))

        def check_common(id_eleve: int | None, id_famille: int | None, date_vers: date | None) -> str:
            if id_eleve is None or id_eleve not in eleve_ids:
                return f"Élève {id_eleve} introuvable dans la nouvelle base"
            if id_famille is None or id_famille not in famille_ids:
                return f"Famille {id_famille} introuvable dans la nouvelle base"
            if date_vers is None:
                return "Date de versement absente ou invalide"
            return ""

        # --- versementscol (scolarité + transport) ---
        for excel_row, source in scol_rows:
            legacy_id = clean(source.get("IDVersementScol"))
            id_eleve = as_int(source.get("IDEleve"))
            id_famille = as_int(source.get("IDFamille"))
            date_vers = as_date(source.get("DateVers"))
            reason = check_common(id_eleve, id_famille, date_vers)
            if reason:
                rejected.append(Rejection("versementscol", excel_row, legacy_id, reason))
                continue
            session.add(VersementScol(
                IDFamille=id_famille,
                IDEleve=id_eleve,
                IDTAnneeScolaire=annee.IDTAnneeScolaire,
                DateVers=date_vers,
                MontantVersSco=as_decimal(source.get("MontantVersSco")),
                MontantVersTrans=as_decimal(source.get("MontantVersTrans")),
                MontantCantine=0,
                MontantVersAutres=0,
                Restitution=as_bool(source.get("Restitution")),
                Reduction=as_bool(source.get("Reduction")),
                Impaye=as_bool(source.get("Impaye")),
                Login=clean(source.get("Login"))[:50] or None,
            ))
            inserted += 1

        # --- versementcant (cantine) ---
        for excel_row, source in cant_rows:
            legacy_id = clean(source.get("IDVersementCant"))
            id_eleve = as_int(source.get("IDEleve"))
            id_famille = as_int(source.get("IDFamille"))
            date_vers = as_date(source.get("DateVers"))
            reason = check_common(id_eleve, id_famille, date_vers)
            if reason:
                rejected.append(Rejection("versementcant", excel_row, legacy_id, reason))
                continue
            session.add(VersementScol(
                IDFamille=id_famille,
                IDEleve=id_eleve,
                IDTAnneeScolaire=annee.IDTAnneeScolaire,
                DateVers=date_vers,
                MontantVersSco=0,
                MontantVersTrans=0,
                MontantCantine=as_decimal(source.get("MontantVers")),
                MontantVersAutres=0,
                Restitution=as_bool(source.get("Restitution")),
                Reduction=as_bool(source.get("Reduction")),
                Impaye=as_bool(source.get("Impaye")),
                Login=clean(source.get("Login"))[:50] or None,
            ))
            inserted += 1

        session.flush()
        if inserted:
            session.execute(text(
                "SELECT setval(pg_get_serial_sequence('\"VersementScol\"', 'IDVersementScol'), "
                "COALESCE((SELECT MAX(\"IDVersementScol\") FROM \"VersementScol\"), 1), true)"
            ))
        write_rejections(args.rejects, rejected)
        if args.commit:
            session.commit()
            action = "IMPORT VALIDÉ"
        else:
            session.rollback()
            action = "SIMULATION : aucune modification enregistrée"
        print(action)
        print(f"Lignes lues : {len(scol_rows) + len(cant_rows)} (scol={len(scol_rows)}, cant={len(cant_rows)})")
        print(f"Lignes prêtes à être insérées : {inserted}")
        print(f"Lignes rejetées/ignorées : {len(rejected)}")
        print(f"Rapport : {args.rejects.resolve()}")
        return 0
    except IntegrityError as exc:
        session.rollback()
        print(f"Import annulé à cause d'une contrainte de base : {exc.orig}", file=sys.stderr)
        return 2
    except Exception as exc:
        session.rollback()
        print(f"Import annulé : {exc}", file=sys.stderr)
        return 2
    finally:
        session.close()


if __name__ == "__main__":
    raise SystemExit(main())
